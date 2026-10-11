from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "brand"
OUT.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUT / "G-Unit-Cakes-Brand-Identity.docx"

LOGO = ROOT / "brand-assets" / "g-unit-cakes-logo-from-screenshot.png"
PACKAGING = ROOT / "static-site" / "assets" / "packaging-system.png"
APPLICATIONS = ROOT / "static-site" / "assets" / "future-use-mockups.png"

INK = "1F1D1D"
CHARCOAL = "2B2928"
CORAL = "EE8D83"
CORAL_DEEP = "C86D65"
BLUSH = "F9D8D4"
CREAM = "FFF8EF"
IVORY = "FFFDF9"
GOLD = "B98642"
KRAFT = "B8895F"
MUTED = "746E6B"
LINE = "E7DDD8"
WHITE = "FFFFFF"

DISPLAY = "Georgia"
BODY = "Arial"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, options in edges.items():
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in options.items():
            element.set(qn(f"w:{key}"), str(value))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_widths(table, widths_inches):
    table.autofit = False
    for row in table.rows:
        for index, width in enumerate(widths_inches):
            row.cells[index].width = Inches(width)
            tc_pr = row.cells[index]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(round(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(round(sum(widths_inches) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")


def set_repeat_no_split(paragraph, keep_with_next=False):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_lines = OxmlElement("w:keepLines")
    p_pr.append(keep_lines)
    if keep_with_next:
        keep_next = OxmlElement("w:keepNext")
        p_pr.append(keep_next)


def set_run(run, font=BODY, size=10.5, color=INK, bold=False, italic=False):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic
    return run


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])


def paragraph_border_bottom(paragraph, color=LINE, size=8, space=8):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    title = styles["Title"]
    title.font.name = DISPLAY
    title._element.rPr.rFonts.set(qn("w:ascii"), DISPLAY)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), DISPLAY)
    title.font.size = Pt(29)
    title.font.bold = False
    title.font.color.rgb = rgb(INK)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    subtitle.font.name = BODY
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), BODY)
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), BODY)
    subtitle.font.size = Pt(11)
    subtitle.font.color.rgb = rgb(MUTED)
    subtitle.paragraph_format.space_after = Pt(10)

    for style_name, size, before, after in (
        ("Heading 1", 22, 0, 9),
        ("Heading 2", 13, 12, 5),
        ("Heading 3", 10.5, 9, 3),
    ):
        style = styles[style_name]
        style.font.name = DISPLAY if style_name != "Heading 3" else BODY
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = style_name == "Heading 3"
        style.font.color.rgb = rgb(INK if style_name == "Heading 1" else CORAL_DEEP)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Kicker" not in styles:
        kicker = styles.add_style("Kicker", WD_STYLE_TYPE.PARAGRAPH)
    else:
        kicker = styles["Kicker"]
    kicker.font.name = BODY
    kicker._element.rPr.rFonts.set(qn("w:ascii"), BODY)
    kicker._element.rPr.rFonts.set(qn("w:hAnsi"), BODY)
    kicker.font.size = Pt(8)
    kicker.font.bold = True
    kicker.font.color.rgb = rgb(CORAL_DEEP)
    kicker.paragraph_format.space_before = Pt(0)
    kicker.paragraph_format.space_after = Pt(5)

    configure_numbering(doc)
    configure_header_footer(section)


def configure_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_id = 42
    num_id = 42
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "260")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "70")
    spacing.set(qn("w:line"), "290")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), CORAL_DEEP)
    r_pr.append(color)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)


def configure_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("G-UNIT CAKES  /  BRAND IDENTITY")
    set_run(r, BODY, 7.5, MUTED, True)
    paragraph_border_bottom(p, LINE, 5, 5)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(3)
    r = p.add_run("VERSION 1.0  /  JULY 2026     ")
    set_run(r, BODY, 7.5, MUTED, True)
    add_field(p, "PAGE")
    paragraph_border_bottom(p, LINE, 4, 4)


def page_break(doc):
    doc.add_page_break()


def add_kicker(doc, text):
    p = doc.add_paragraph(style="Kicker")
    p.add_run(text.upper())
    return p


def add_page_title(doc, number, title, subtitle=None):
    add_kicker(doc, f"{number}  /  G-UNIT CAKES")
    p = doc.add_paragraph(title, style="Heading 1")
    set_repeat_no_split(p, True)
    if subtitle:
        s = doc.add_paragraph(subtitle, style="Subtitle")
        set_repeat_no_split(s)


def add_lead(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.25
    set_run(p.add_run(text), DISPLAY, 14, INK)
    return p


def add_body(doc, text, bold_label=None):
    p = doc.add_paragraph()
    if bold_label:
        set_run(p.add_run(bold_label), BODY, 10.5, INK, True)
    set_run(p.add_run(text), BODY, 10.5, INK)
    return p


def add_bullet(doc, text, bold_label=None):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "42")
    num_pr.extend([ilvl, num_id])
    p_pr.append(num_pr)
    if bold_label:
        set_run(p.add_run(bold_label), BODY, 10.2, INK, True)
    set_run(p.add_run(text), BODY, 10.2, INK)
    return p


def add_callout(doc, label, text, fill=CREAM, accent=CORAL_DEEP):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_widths(table, [6.45])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=180, start=220, bottom=180, end=220)
    set_cell_border(
        cell,
        top={"val": "single", "sz": "6", "color": fill},
        bottom={"val": "single", "sz": "6", "color": fill},
        start={"val": "single", "sz": "16", "color": accent},
        end={"val": "single", "sz": "6", "color": fill},
    )
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run(label.upper()), BODY, 8, accent, True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    set_run(p2.add_run(text), DISPLAY, 13, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_two_column_items(doc, items):
    table = doc.add_table(rows=(len(items) + 1) // 2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [3.18, 3.18])
    for index, (title, detail) in enumerate(items):
        row = index // 2
        col = index % 2
        cell = table.cell(row, col)
        set_cell_margins(cell, top=130, start=150, bottom=150, end=170)
        set_cell_border(
            cell,
            bottom={"val": "single", "sz": "5", "color": LINE},
        )
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        set_run(p.add_run(title), BODY, 9.5, CORAL_DEEP, True)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        set_run(p2.add_run(detail), BODY, 9.5, INK)
    return table


def add_palette_table(doc, colors):
    table = doc.add_table(rows=len(colors), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [1.08, 1.9, 3.38])
    for row, (name, hex_value, rgb_value, role) in zip(table.rows, colors):
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(row.cells[0], hex_value)
        for cell in row.cells:
            set_cell_margins(cell, top=110, start=130, bottom=110, end=130)
            set_cell_border(
                cell,
                top={"val": "single", "sz": "4", "color": LINE},
                bottom={"val": "single", "sz": "4", "color": LINE},
                start={"val": "single", "sz": "4", "color": LINE},
                end={"val": "single", "sz": "4", "color": LINE},
            )
        p = row.cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        set_run(p.add_run(name), BODY, 9.5, INK, True)
        p2 = row.cells[1].add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        set_run(p2.add_run(f"#{hex_value}  |  {rgb_value}"), BODY, 8.5, MUTED)
        p3 = row.cells[2].paragraphs[0]
        p3.paragraph_format.space_after = Pt(0)
        set_run(p3.add_run(role), BODY, 9.2, INK)
    return table


def add_logo_box(doc, image_path, width=2.6, fill=IVORY):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [6.45])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=220, start=220, bottom=220, end=220)
    set_cell_border(
        cell,
        top={"val": "single", "sz": "5", "color": LINE},
        bottom={"val": "single", "sz": "5", "color": LINE},
        start={"val": "single", "sz": "5", "color": LINE},
        end={"val": "single", "sz": "5", "color": LINE},
    )
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    return table


def build_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("G-UNIT CAKES"), BODY, 8.5, CORAL_DEEP, True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO), width=Inches(3.15))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    set_run(p.add_run("Brand Identity"), DISPLAY, 31, INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    set_run(
        p.add_run("Visual identity, voice, packaging, and digital guidance"),
        BODY,
        10.5,
        MUTED,
    )

    add_callout(doc, "Brand essence", "Celebration, made personal.", BLUSH, CORAL_DEEP)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    set_run(p.add_run("VERSION 1.0  /  JULY 2026"), BODY, 8, MUTED, True)


def build_brand_core(doc):
    page_break(doc)
    add_page_title(
        doc,
        "01",
        "Brand core",
        "The strategic centre that keeps every cake, message, package, and customer interaction recognisably G-Unit Cakes.",
    )
    add_lead(doc, "A warm, elegant cake studio built around personal celebration.")
    add_callout(
        doc,
        "Brand promise",
        "Every order should feel thoughtfully designed, carefully made, clearly communicated, and joyful to receive.",
    )
    doc.add_paragraph("Foundation", style="Heading 2")
    add_two_column_items(
        doc,
        [
            (
                "Mission",
                "To bake memorable cakes and sweet details that make every celebration feel personal.",
            ),
            (
                "Vision",
                "To become a trusted local cake studio known for thoughtful design, generous flavour, and warm service.",
            ),
            (
                "Positioning",
                "A custom cake studio for people who want a celebration piece that looks considered, tastes generous, and feels made specifically for their moment.",
            ),
            (
                "Tagline",
                "Made for your moment.",
            ),
        ],
    )
    doc.add_paragraph("Personality", style="Heading 2")
    for label, detail in (
        ("Elegant, not formal. ", "Polished enough for a milestone, welcoming enough for every day."),
        ("Warm, not overly familiar. ", "Personal care without forced intimacy."),
        ("Creative, not chaotic. ", "Expressive design guided by clear choices."),
        ("Confident, not boastful. ", "Calm expertise that helps customers decide."),
        ("Reliable, not transactional. ", "Organised service from first enquiry to handover."),
    ):
        add_bullet(doc, detail, label)


def build_audience_messages(doc):
    page_break(doc)
    add_page_title(
        doc,
        "02",
        "Audience and message",
        "Speak first to the meaning of the occasion, then to design, flavour, and a clear ordering experience.",
    )
    add_lead(doc, "Beautifully considered cakes, baked for the moments worth celebrating.")
    doc.add_paragraph("Who we serve", style="Heading 2")
    add_two_column_items(
        doc,
        [
            (
                "Celebration customers",
                "Birthdays, graduations, baby showers, anniversaries, weddings, and personal milestones.",
            ),
            (
                "Thoughtful gift buyers",
                "People sending a cake, dessert box, cupcakes, or cookies to make someone feel seen.",
            ),
            (
                "Event partners",
                "Planners and hosts who need a dependable centrepiece that fits the wider event.",
            ),
            (
                "Small businesses",
                "Teams ordering branded treats, thank-you gifts, launches, or client celebrations.",
            ),
        ],
    )
    doc.add_paragraph("Message pillars", style="Heading 2")
    for label, detail in (
        ("Personal by design. ", "Every order begins with the occasion, taste, and visual references."),
        ("Beautiful and delicious. ", "The finish creates a strong first impression without compromising flavour."),
        ("Clear from enquiry to collection. ", "Ordering feels welcoming, organised, and dependable."),
        ("Thoughtful in every detail. ", "Decoration, finishing, packaging, and handover feel connected."),
    ):
        add_bullet(doc, detail, label)
    doc.add_paragraph("Preferred calls to action", style="Heading 2")
    add_callout(
        doc,
        "Use direct language",
        "Start your custom order  /  Tell us about your celebration  /  Check availability  /  Order on WhatsApp",
        CREAM,
        GOLD,
    )


def build_voice(doc):
    page_break(doc)
    add_page_title(
        doc,
        "03",
        "Voice and tone",
        "The brand sounds like a skilled baker who is warm, calm, creative, and easy to order from.",
    )
    add_two_column_items(
        doc,
        [
            ("Warm", "Use genuine care and hospitality. Acknowledge the occasion."),
            ("Celebratory", "Bring energy to the moment without filling every line with exclamation marks."),
            ("Clear", "Make dates, portions, pricing, and next steps easy to understand."),
            ("Confident", "Guide customers with calm expertise and useful recommendations."),
            ("Personal", "Sound like a real baker speaking to a real customer."),
            ("Concise", "Use short sentences and specific details."),
        ],
    )
    doc.add_paragraph("Writing pattern", style="Heading 2")
    for text in (
        "Start with the customer's occasion or desired outcome.",
        "Add one appetising or visual detail.",
        "State the useful information clearly.",
        "End with one simple next step.",
    ):
        add_bullet(doc, text)
    doc.add_paragraph("Social caption example", style="Heading 2")
    add_callout(
        doc,
        "Caption",
        "Soft coral piping, hand-finished details, and a centre made for celebrating. Planning something special? Tell us the date, servings, flavours, and the look you love.",
        BLUSH,
        CORAL_DEEP,
    )
    doc.add_paragraph("Enquiry reply example", style="Heading 2")
    add_body(
        doc,
        "Thank you for thinking of G-Unit Cakes for your celebration. Please send your event date, number of servings, preferred flavours, and any inspiration images. We will confirm availability and guide you through the next step.",
    )


def build_logo(doc):
    page_break(doc)
    add_page_title(
        doc,
        "04",
        "Primary logo",
        "The whisk-shaped G is the recognisable heart of the identity. The full lockup is the default brand signature.",
    )
    add_logo_box(doc, LOGO, 2.45)
    doc.add_paragraph("Anatomy and meaning", style="Heading 2")
    add_two_column_items(
        doc,
        [
            ("Whisk", "Craft, movement, and handmade baking."),
            ("Letter G", "Immediate recognition and ownership of the name."),
            ("Frosting flourish", "Celebration, decoration, and expressive cake design."),
            ("Heart detail", "Care, warmth, and personal service."),
            ("Gold handle", "A crafted, warm, quality accent."),
            ("Serif wordmark", "A refined visual tone with a timeless feel."),
        ],
    )
    add_callout(
        doc,
        "Default choice",
        "Use the full-colour primary logo on white, ivory, blush, or another pale, quiet background.",
        CREAM,
        GOLD,
    )


def build_logo_rules(doc):
    page_break(doc)
    add_page_title(
        doc,
        "05",
        "Logo use",
        "Protect the mark with enough space, enough size, and a background that keeps every detail legible.",
    )
    doc.add_paragraph("Clear space", style="Heading 2")
    add_logo_box(doc, LOGO, 1.7, CREAM)
    add_body(
        doc,
        "Keep clear space around the logo equal to the height of the word CAKES. No text, border, image, or package edge should enter this zone.",
    )
    doc.add_paragraph("Minimum size", style="Heading 2")
    add_two_column_items(
        doc,
        [
            ("Digital - full logo", "140 px wide minimum"),
            ("Digital - symbol", "40 px wide minimum"),
            ("Print - full logo", "38 mm wide minimum"),
            ("Print - symbol", "12 mm wide minimum"),
        ],
    )
    doc.add_paragraph("Never", style="Heading 2")
    for text in (
        "Stretch, compress, rotate, crop, or rearrange the logo.",
        "Change individual logo colours or add unapproved colours.",
        "Add shadows, bevels, glows, outlines, transparency, or gradients.",
        "Place the logo over busy cake photography without a quiet field.",
        "Use the detailed full logo below the minimum size.",
    ):
        add_bullet(doc, text)
    add_callout(
        doc,
        "Production note",
        "The screenshot-matched SVG embeds cleaned raster artwork. Before signage, embroidery, foil stamping, laser cutting, or large-format print, approve a fully path-traced vector master.",
        "FFF2E4",
        GOLD,
    )


def build_colour(doc):
    page_break(doc)
    add_page_title(
        doc,
        "06",
        "Colour palette",
        "A light, appetite-led system: quiet ivory fields, strong charcoal structure, coral warmth, and restrained gold detail.",
    )
    colors = [
        ("Signature Coral", CORAL, "238, 141, 131", "Frosting, highlights, social accents"),
        ("Charcoal", INK, "31, 29, 29", "Wordmark, body text, premium contrast"),
        ("Ivory", IVORY, "255, 253, 249", "Primary background"),
        ("Deep Coral", CORAL_DEEP, "200, 109, 101", "Links, buttons, accessible emphasis"),
        ("Soft Blush", BLUSH, "249, 216, 212", "Packaging and gentle background panels"),
        ("Whisk Gold", GOLD, "185, 134, 66", "Small premium accents"),
        ("Kraft", KRAFT, "184, 137, 95", "Natural packaging and texture"),
    ]
    add_palette_table(doc, colors)
    doc.add_paragraph("Recommended balance", style="Heading 2")
    add_body(
        doc,
        "Use ivory or white for most of the field, charcoal for structure and readability, coral as the main emotional accent, and gold sparingly. Kraft and blush support packaging and photography.",
    )
    add_callout(
        doc,
        "Accessibility",
        "Use charcoal for normal text. Deep coral can be used for larger headings and controls; test every text/background pairing to WCAG AA before launch.",
        CREAM,
        CORAL_DEEP,
    )


def build_typography(doc):
    page_break(doc)
    add_page_title(
        doc,
        "07",
        "Typography",
        "Pair a graceful editorial serif with a direct, modern sans serif. The contrast mirrors decorated craft and dependable service.",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run("Cormorant Garamond"), DISPLAY, 25, INK)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(15)
    set_run(
        p2.add_run("Preferred display face  /  Medium or Semibold  /  Headings, quotes, campaign moments"),
        BODY,
        9.3,
        MUTED,
    )
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_after = Pt(2)
    set_run(p3.add_run("Manrope"), BODY, 23, INK, True)
    p4 = doc.add_paragraph()
    p4.paragraph_format.space_after = Pt(15)
    set_run(
        p4.add_run("Preferred body face  /  Regular or Medium  /  Website, forms, captions, order information"),
        BODY,
        9.3,
        MUTED,
    )
    doc.add_paragraph("Reliable fallbacks", style="Heading 2")
    add_two_column_items(
        doc,
        [
            ("Display", "Georgia, Times New Roman, serif"),
            ("Body", "Arial, Helvetica, sans-serif"),
        ],
    )
    doc.add_paragraph("Type rules", style="Heading 2")
    for text in (
        "Use sentence case for headings and calls to action.",
        "Reserve all caps for short labels only.",
        "Use body text at 16 px minimum in digital applications.",
        "Keep letter spacing at 0 for headings and body copy.",
        "Do not use more than two font families in one design.",
        "Keep line length comfortable: roughly 55 to 75 characters.",
    ):
        add_bullet(doc, text)
    add_callout(
        doc,
        "Simple hierarchy",
        "Elegant display headline  /  Clear supporting line  /  Direct action in the sans serif",
        BLUSH,
        CORAL_DEEP,
    )


def build_visual_language(doc):
    page_break(doc)
    add_page_title(
        doc,
        "08",
        "Visual language",
        "The cake remains the hero. Brand details provide recognition, order, and warmth around it.",
    )
    doc.add_paragraph("Photography", style="Heading 2")
    for label, detail in (
        ("Light. ", "Use bright natural light and keep frosting colours accurate."),
        ("Subject. ", "Show the real cake clearly before atmospheric details."),
        ("Variety. ", "Mix complete cake portraits, close piping details, slices, and hands at work."),
        ("Surface. ", "Prefer ivory, blush, stone, pale wood, charcoal, or natural kraft."),
        ("Edit. ", "Keep the result warm, soft, and true to the product."),
    ):
        add_bullet(doc, detail, label)
    doc.add_paragraph("Composition", style="Heading 2")
    add_two_column_items(
        doc,
        [
            ("Quiet space", "Leave room for short copy without covering the cake."),
            ("Product clarity", "Use straight-on or gentle three-quarter views for hero images."),
            ("Purposeful crop", "Use close detail crops deliberately; never clip the cake accidentally."),
            ("Relevant props", "Use only a few props tied to flavour, occasion, or making."),
        ],
    )
    doc.add_paragraph("Graphic devices", style="Heading 2")
    add_callout(
        doc,
        "Recognisable details",
        "Fine coral rules and flourishes  /  Small heart motif  /  Circular logo seals  /  Ivory fields  /  Blush or coral package bands",
        CREAM,
        GOLD,
    )


def build_packaging(doc):
    page_break(doc)
    add_page_title(
        doc,
        "09",
        "Packaging system",
        "A flexible family of white, blush, kraft, ivory, and coral materials that feels coherent at every order size.",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    p.add_run().add_picture(str(PACKAGING), width=Inches(4.95))
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    set_run(
        p2.add_run("Cake box  /  handled box  /  cupcake sleeve  /  thank-you card  /  labels  /  treat bag"),
        BODY,
        8.5,
        MUTED,
    )


def build_applications(doc):
    page_break(doc)
    add_page_title(
        doc,
        "10",
        "Brand applications",
        "Use the same visual cues at the moments customers see, carry, photograph, gift, and share the product.",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    p.add_run().add_picture(str(APPLICATIONS), width=Inches(4.95))
    add_callout(
        doc,
        "Application principle",
        "Use one confident logo placement, one clear message, and enough quiet space for the cake or product to remain the focus.",
        CREAM,
        CORAL_DEEP,
    )


def build_digital(doc):
    page_break(doc)
    add_page_title(
        doc,
        "11",
        "Digital and social",
        "Digital touchpoints should make discovery appetising and ordering reassuringly simple.",
    )
    doc.add_paragraph("Website", style="Heading 2")
    for text in (
        "Use the full logo in the header and a simplified symbol for the favicon.",
        "Lead with clear cake photography and a direct custom-order action.",
        "Collect event date, servings, flavour, style, budget, delivery or collection, and inspiration images.",
        "Keep Instagram, WhatsApp, and TikTok links as placeholders until the final accounts are supplied.",
    ):
        add_bullet(doc, text)
    doc.add_paragraph("Social formats", style="Heading 2")
    add_two_column_items(
        doc,
        [
            ("Feed posts", "4:5 portrait. Keep text overlays short and inside safe margins."),
            ("Stories and video", "9:16. Start with the finished cake or a strong making moment."),
            ("Profile image", "Use the circular G/whisk seal once a clean symbol-only asset is approved."),
            ("Content rhythm", "Alternate finished cakes, details, process, customer moments, and ordering prompts."),
        ],
    )
    doc.add_paragraph("WhatsApp ordering", style="Heading 2")
    add_callout(
        doc,
        "Order summary",
        "Design  /  servings  /  flavour  /  price  /  payment status  /  collection or delivery  /  final date and time",
        BLUSH,
        CORAL_DEEP,
    )
    add_body(
        doc,
        "Confirm availability before detailed decoration discussions, then keep the final agreed details together in one clear message.",
    )


def build_asset_control(doc):
    page_break(doc)
    add_page_title(
        doc,
        "12",
        "Asset control",
        "Consistent files and approvals protect the identity as the bakery grows.",
    )
    doc.add_paragraph("Current logo assets", style="Heading 2")
    assets = [
        ("Screenshot-matched SVG", "g-unit-cakes-logo-from-screenshot.svg"),
        ("Transparent PNG", "g-unit-cakes-logo-from-screenshot.png"),
        ("Editable vector recreation", "g-unit-cakes-logo.svg"),
        ("High-resolution PNG", "g-unit-cakes-logo.png"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [2.05, 4.31])
    headers = table.rows[0].cells
    for cell, text in zip(headers, ("Asset", "File")):
        set_cell_shading(cell, BLUSH)
        set_cell_margins(cell)
        set_run(cell.paragraphs[0].add_run(text), BODY, 9, INK, True)
    set_repeat_table_header(table.rows[0])
    for label, filename in assets:
        row = table.add_row().cells
        for cell in row:
            set_cell_margins(cell)
            set_cell_border(
                cell,
                bottom={"val": "single", "sz": "4", "color": LINE},
            )
        set_run(row[0].paragraphs[0].add_run(label), BODY, 9.2, INK, True)
        set_run(row[1].paragraphs[0].add_run(filename), BODY, 8.8, MUTED)
    doc.add_paragraph("Naming convention", style="Heading 2")
    add_callout(
        doc,
        "File pattern",
        "g-unit-cakes-[asset]-[variant]-[colour]-[size].[ext]",
        CREAM,
        GOLD,
    )
    add_body(
        doc,
        "Example: g-unit-cakes-logo-stacked-full-colour-2000px.png",
    )
    doc.add_paragraph("Approval checklist", style="Heading 2")
    for text in (
        "Correct logo version, colour, clear space, and minimum size.",
        "Approved typefaces and palette.",
        "Cake is clear, accurate, and appetising.",
        "Copy sounds warm, clear, and celebratory.",
        "Order details, prices, dates, and contact links are correct.",
        "Final artwork is proofed in its actual format before publishing or printing.",
    ):
        add_bullet(doc, text)


def set_document_properties(doc):
    props = doc.core_properties
    props.title = "G-Unit Cakes Brand Identity"
    props.subject = "Visual identity, voice, packaging, and digital guidelines"
    props.author = "G-Unit Cakes"
    props.keywords = "G-Unit Cakes, brand identity, logo, packaging, bakery"
    props.comments = "Version 1.0 - July 2026"


def main():
    doc = Document()
    configure_document(doc)
    set_document_properties(doc)
    build_cover(doc)
    build_brand_core(doc)
    build_audience_messages(doc)
    build_voice(doc)
    build_logo(doc)
    build_logo_rules(doc)
    build_colour(doc)
    build_typography(doc)
    build_visual_language(doc)
    build_packaging(doc)
    build_applications(doc)
    build_digital(doc)
    build_asset_control(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
